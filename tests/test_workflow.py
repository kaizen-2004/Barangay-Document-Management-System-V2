import os
from datetime import date

from docx import Document as DocxDocument
from PIL import Image

from barangay_project.extensions import db
from barangay_project.models import Document, Official
from barangay_project.docx_pipeline import resolve_stored_path_to_abs


def _login(client, username, password):
    return client.post(
        "/login",
        data={"username": username, "password": password},
        follow_redirects=False,
    )


def test_document_workflow(client, app, make_user, make_resident, make_document_type):
    make_user("clerk", "Clerk123!", role="clerk")
    make_user("boss", "Admin123!", role="admin")

    resident = make_resident(birth_date=date(1990, 1, 1))
    doc_type = make_document_type(name="Test Clearance", requires_photo=False)

    template_dir = app.config["DOCX_TEMPLATE_UPLOAD_DIR"]
    os.makedirs(template_dir, exist_ok=True)
    os.makedirs(os.path.join(app.config["UPLOAD_FOLDER"], "residents"), exist_ok=True)

    template_abs = os.path.join(template_dir, "test-clearance.docx")
    tpl = DocxDocument()
    tpl.add_paragraph("Resident: {{ resident_name }}")
    tpl.add_paragraph("Address: {{ address }}")
    tpl.add_paragraph("Purpose: {{ purpose }}")
    tpl.add_paragraph("Issue: {{ issue_date }}")
    tpl.add_paragraph("Photo: {{ resident_photo }}")
    tpl.add_paragraph("Captain: {{ captain_signature }}")
    tpl.add_paragraph("QR: {{ qr_code }}")
    tpl.save(template_abs)

    resident_photo_abs = os.path.join(app.config["UPLOAD_FOLDER"], "residents", "resident-photo.jpg")
    Image.new("RGB", (250, 250), color=(200, 200, 200)).save(resident_photo_abs)

    resident.photo_path = "uploads/residents/resident-photo.jpg"
    doc_type.template_path = "test-clearance.docx"
    doc_type.template_filename = "test-clearance.docx"
    doc_type.template_active = True

    official = Official(
        full_name="Juan Dela Cruz",
        title="Barangay Captain",
        is_active=True,
    )
    db.session.add(official)
    db.session.commit()

    _login(client, "clerk", "Clerk123!")

    resp = client.post(
        "/documents/issue",
        data={
            "resident_id": resident.id,
            "document_type_id": doc_type.id,
            "details": "Test document",
            "issue_date": date.today().isoformat(),
        },
        follow_redirects=False,
    )
    assert resp.status_code == 302

    doc = Document.query.first()
    assert doc is not None
    assert doc.status == "draft"

    resp = client.post(f"/documents/{doc.id}/request-approval", follow_redirects=False)
    assert resp.status_code == 302
    doc = db.session.get(Document, doc.id)
    assert doc.status == "pending"

    client.get("/logout", follow_redirects=False)
    _login(client, "boss", "Admin123!")

    resp = client.post(f"/documents/{doc.id}/approve", follow_redirects=False)
    assert resp.status_code == 302
    doc = db.session.get(Document, doc.id)
    assert doc.status == "approved"

    resp = client.post(f"/documents/{doc.id}/issue", follow_redirects=False)
    assert resp.status_code == 302
    doc = db.session.get(Document, doc.id)
    assert doc.status == "issued"
    assert doc.file_path

    rel = doc.file_path
    abs_path = resolve_stored_path_to_abs(rel)
    assert abs_path is not None
    assert os.path.exists(abs_path)
